#!/usr/bin/env python3
"""
Unlock and extract text from the password-protected documents Indian taxpayers
actually receive: Form 16 / AIS / 26AS / ITR PDFs, bank interest certificates,
and encrypted Excel statements (PPF/CAS/broker exports saved as OLE2 containers).

Tries a candidate-password list built from:
  - passwords given on the command line (--passwords)
  - PAN + DOB combos (the portal's standard: lowercase PAN + DOB ddmmyyyy)
  - digit/alphanumeric tokens mined from the FILENAME (users often save files
    as "statement - <password>.pdf")
  - the empty password (some "protected" PDFs open with owner-password only)

Output: <name>_extracted.txt next to each input file (PDF text or sheet dump),
plus a per-file report of which password worked.

Usage:
    python unlock_documents.py <file-or-directory> [--pan ABCPD1234E] [--dob 15061990]
                               [--passwords pw1 pw2 ...] [--no-write]

Dependencies: pip install pypdf msoffcrypto-tool openpyxl xlrd
(only what a given file type needs is imported at runtime)
"""

import argparse
import io
import re
import sys
from pathlib import Path


def candidate_passwords(filename, pan=None, dob=None, extra=None):
    cands, seen = [], set()

    def add(p):
        if p is not None and p not in seen:
            seen.add(p)
            cands.append(p)

    for p in extra or []:
        add(p)
    if pan and dob:
        add(pan.lower() + dob)
        add(pan.upper() + dob)
    if pan:
        add(pan.lower())
        add(pan.upper())
    if dob:
        add(dob)
    # tokens mined from the filename: runs of >=4 alphanumerics
    stem = Path(filename).stem
    for tok in re.findall(r"[A-Za-z0-9]{4,}", stem):
        add(tok)
        add(tok.lower())
    # pairwise joins of filename tokens (e.g. "abcpd1234e" + "15061990")
    toks = re.findall(r"[A-Za-z0-9]{4,}", stem)
    for a in toks:
        for b in toks:
            if a != b:
                add(a.lower() + b)
    add("")  # owner-password-only PDFs
    return cands


def extract_pdf(path, passwords):
    from pypdf import PdfReader

    for pwd in passwords:
        try:
            reader = PdfReader(path)
            if reader.is_encrypted:
                if not reader.decrypt(pwd):
                    continue
            text = "\n".join(
                f"===== PAGE {i + 1} =====\n{(pg.extract_text() or '').strip()}"
                for i, pg in enumerate(reader.pages)
            )
            return pwd, text
        except Exception:
            continue
    return None, None


def _dump_openpyxl(buf_or_path):
    import openpyxl

    wb = openpyxl.load_workbook(buf_or_path, data_only=True)
    out = []
    for sn in wb.sheetnames:
        out.append(f"===== SHEET {sn} =====")
        for row in wb[sn].iter_rows():
            vals = [str(c.value) for c in row if c.value is not None and str(c.value).strip()]
            if vals:
                out.append("\t".join(vals))
    return "\n".join(out)


def _dump_xlrd(buf_bytes):
    import xlrd

    wb = xlrd.open_workbook(file_contents=buf_bytes)
    out = []
    for sn in wb.sheet_names():
        sh = wb.sheet_by_name(sn)
        out.append(f"===== SHEET {sn} =====")
        for r in range(sh.nrows):
            vals = [str(sh.cell_value(r, c)) for c in range(sh.ncols) if str(sh.cell_value(r, c)).strip()]
            if vals:
                out.append("\t".join(vals))
    return "\n".join(out)


def extract_excel(path, passwords):
    """Handles plain xlsx, legacy xls, and OLE2-encrypted workbooks
    (a '.xlsx' that openpyxl rejects as 'not a zip file' is usually an
    encrypted OLE2 container -> msoffcrypto with a password)."""
    raw = Path(path).read_bytes()
    # 1) plain xlsx (zip magic)
    if raw[:4] == b"PK\x03\x04":
        return "", _dump_openpyxl(path)
    # 2) OLE2 container: encrypted xlsx OR legacy xls
    if raw[:8] == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1":
        try:
            return "", _dump_xlrd(raw)  # genuine legacy xls, no password
        except Exception:
            pass
        import msoffcrypto

        for pwd in passwords:
            if pwd == "":
                continue
            try:
                with open(path, "rb") as f:
                    of = msoffcrypto.OfficeFile(f)
                    of.load_key(password=pwd)
                    buf = io.BytesIO()
                    of.decrypt(buf)
                buf.seek(0)
                try:
                    return pwd, _dump_openpyxl(buf)
                except Exception:
                    return pwd, _dump_xlrd(buf.getvalue())
            except Exception:
                continue
    return None, None


def process(path, pan, dob, extra, write=True):
    path = Path(path)
    passwords = candidate_passwords(path.name, pan, dob, extra)
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        pwd, text = extract_pdf(path, passwords)
    elif suffix in (".xlsx", ".xls", ".xlsm"):
        pwd, text = extract_excel(path, passwords)
    else:
        print(f"SKIP  {path.name} (unsupported type)")
        return
    if text is None:
        print(f"FAIL  {path.name} — no candidate password worked; ask the user for it")
        return
    label = "(no password)" if pwd == "" else f"(password: {pwd})"
    print(f"OK    {path.name} {label}")
    if write:
        out = path.with_name(path.stem + "_extracted.txt")
        out.write_text(text, encoding="utf-8")
        print(f"      -> {out.name}")
    else:
        print(text)


if __name__ == "__main__":
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("target", help="A document file, or a directory to sweep")
    ap.add_argument("--pan", default=None)
    ap.add_argument("--dob", default=None, help="ddmmyyyy")
    ap.add_argument("--passwords", nargs="*", default=[], help="Known passwords to try first")
    ap.add_argument("--no-write", action="store_true", help="Print extracted text to stdout instead of files")
    args = ap.parse_args()

    target = Path(args.target)
    files = sorted(p for p in target.iterdir() if p.is_file()) if target.is_dir() else [target]
    for f in files:
        if f.suffix.lower() in (".pdf", ".xlsx", ".xls", ".xlsm"):
            process(f, args.pan, args.dob, args.passwords, write=not args.no_write)
